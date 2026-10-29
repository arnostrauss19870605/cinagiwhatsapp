"""Bring people into an audience from a spreadsheet the user made themselves.

Two halves, so the browser can show the columns before anything is written:
`read_sheet` turns an upload into plain rows of text, and `import_rows` applies
a column mapping to those rows. A person is matched on their mobile number -
the only thing a spreadsheet and WhatsApp reliably agree on.
"""

import csv
import io
from dataclasses import dataclass, field

from django.db import transaction
from django.utils import timezone

from .models import Contact, normalise_msisdn

MAX_ROWS = 10_000

# The things a column can be. Order is the order the mapping form shows them.
FIELDS = (
    ("mobile", "Mobile number"),
    ("full_name", "Full name"),
    ("first_name", "Name"),
    ("last_name", "Surname"),
)

# Lower-cased heading words that probably mean each field.
_GUESSES = {
    "mobile": ("mobile", "cell", "cellphone", "cell number", "cell no", "mobile number", "phone",
               "phone number", "msisdn", "whatsapp", "whatsapp number", "number", "contact number", "tel"),
    "full_name": ("full name", "full_name", "fullname", "name and surname", "contact name", "display name"),
    "first_name": ("name", "first name", "first_name", "firstname", "given name", "forename", "naam"),
    "last_name": ("surname", "last name", "last_name", "lastname", "family name", "van"),
}


class SheetError(ValueError):
    """The upload could not be read. The message is shown to the user as is."""


def _text(value):
    if value is None:
        return ""
    # Excel stores 0821234567 as the number 821234567.0; the float must not
    # leave a trailing ".0" digit behind for the number cleaner to pick up.
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _rows_from_xlsx(upload):
    from openpyxl import load_workbook

    try:
        book = load_workbook(upload, read_only=True, data_only=True)
    except Exception as exc:
        raise SheetError(
            "That file could not be opened as an Excel workbook. Save it as .xlsx and try again."
        ) from exc
    try:
        sheet = book.worksheets[0]
        for row in sheet.iter_rows(values_only=True):
            yield [_text(value) for value in row]
    finally:
        book.close()


def _rows_from_csv(upload):
    raw = upload.read()
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise SheetError("That file is not readable as a CSV. Save it as .xlsx or CSV and try again.")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    for row in csv.reader(io.StringIO(text), dialect):
        yield [_text(value) for value in row]


def read_sheet(upload, *, has_headings=True):
    """(column headings, data rows) from an uploaded .xlsx or .csv.

    Blank rows are dropped and every row is padded to the same width, so a
    column index means the same thing on every row.
    """
    name = (getattr(upload, "name", "") or "").lower()
    if name.endswith(".xlsx") or name.endswith(".xlsm"):
        source = _rows_from_xlsx(upload)
    elif name.endswith(".csv") or name.endswith(".txt"):
        source = _rows_from_csv(upload)
    elif name.endswith(".xls"):
        raise SheetError(
            "That is an old-style Excel file (.xls). Open it in Excel, choose Save As, "
            "pick 'Excel Workbook (.xlsx)' and upload that instead."
        )
    else:
        raise SheetError("Upload an Excel file (.xlsx) or a CSV file.")

    rows = []
    for row in source:
        if not any(row):
            continue
        rows.append(row)
        if len(rows) > MAX_ROWS + 1:
            raise SheetError(
                f"That sheet has more than {MAX_ROWS:,} people in it. Split it into smaller files and upload them one at a time."
            )
    if not rows:
        raise SheetError("That sheet is empty.")

    width = max(len(row) for row in rows)
    while width and not any(len(row) >= width and row[width - 1] for row in rows):
        width -= 1
    rows = [(row + [""] * width)[:width] for row in rows]

    if has_headings:
        headings, rows = rows[0], rows[1:]
        headings = [h or f"Column {_letter(i)}" for i, h in enumerate(headings)]
    else:
        headings = [f"Column {_letter(i)}" for i in range(width)]
    if not rows:
        raise SheetError("That sheet only has a heading row - there is nobody in it to import.")
    return headings, rows


def _letter(index):
    letters = ""
    index += 1
    while index:
        index, remainder = divmod(index - 1, 26)
        letters = chr(65 + remainder) + letters
    return letters


def guess_mapping(headings):
    """{field: column index or None}, from the heading wording."""
    lowered = [" ".join(h.lower().replace("_", " ").split()) for h in headings]
    mapping, taken = {}, set()
    for key, _label in FIELDS:
        mapping[key] = None
        for guess in _GUESSES[key]:
            guess = guess.replace("_", " ")
            if guess in lowered and lowered.index(guess) not in taken:
                mapping[key] = lowered.index(guess)
                taken.add(mapping[key])
                break
    if mapping["full_name"] is not None and mapping["first_name"] is None and mapping["last_name"] is None:
        return mapping
    if mapping["first_name"] is not None and mapping["last_name"] is None and mapping["full_name"] is None:
        # A lone "Name" column is almost always the whole name.
        mapping["full_name"], mapping["first_name"] = mapping["first_name"], None
    return mapping


@dataclass
class ImportResult:
    added: int = 0  # put into the audience, whether or not they were already a contact
    created: int = 0  # of those, brand-new contacts
    updated: list = field(default_factory=list)  # (number, old name, new name) for anyone renamed
    unchanged: int = 0  # already in the audience under the same name
    repeated: int = 0  # the same number further down the sheet
    unusable: list = field(default_factory=list)  # (sheet row number, what was in the cell)


def _cell(row, index):
    if index is None or index >= len(row):
        return ""
    return " ".join(row[index].split())


def _chunks(items, size=500):
    items = list(items)
    for start in range(0, len(items), size):
        yield items[start:start + size]


def import_rows(workspace, audience, rows, mapping, *, first_row_number=2):
    """Put the people in `rows` into `audience`.

    Somebody already in the audience is left alone unless the sheet gives them
    a different name, in which case the name is updated. Somebody who is a
    contact but not in the audience is added (and renamed if the sheet says
    so). Anyone else becomes a new contact. A number that appears twice in the
    sheet is only taken the first time.
    """
    result = ImportResult()
    wanted = {}  # msisdn -> (full name, first name, surname)
    for offset, row in enumerate(rows):
        raw_number = _cell(row, mapping.get("mobile"))
        msisdn = normalise_msisdn(raw_number)
        if len(msisdn) < 10 or len(msisdn) > 15:
            result.unusable.append((first_row_number + offset, raw_number))
            continue
        if msisdn in wanted:
            result.repeated += 1
            continue
        first, last = _cell(row, mapping.get("first_name")), _cell(row, mapping.get("last_name"))
        full = _cell(row, mapping.get("full_name")) or " ".join(part for part in (first, last) if part)
        wanted[msisdn] = (full, first, last)

    with transaction.atomic():
        existing = {}
        for chunk in _chunks(wanted):
            for contact in Contact.objects.for_workspace(workspace).filter(wa_id__in=chunk):
                existing[contact.wa_id] = contact
        members = set(audience.contacts.values_list("pk", flat=True))

        now = timezone.now()
        to_create, to_update, to_add = [], [], []
        for msisdn, (full, first, last) in wanted.items():
            full = full[:120]
            details = {k: v for k, v in (("first_name", first), ("last_name", last)) if v}
            contact = existing.get(msisdn)
            if contact is None:
                to_create.append(Contact(
                    workspace=workspace, wa_id=msisdn, display_name=full, attributes=details,
                ))
                continue

            renamed = bool(full) and contact.display_name != full
            attributes = dict(contact.attributes or {})
            if renamed or any(attributes.get(k) != v for k, v in details.items()):
                if renamed:
                    result.updated.append((contact.pretty_number, contact.name, full))
                    contact.display_name = full
                attributes.update(details)
                contact.attributes = attributes
                contact.updated_at = now
                to_update.append(contact)

            if contact.pk not in members:
                to_add.append(contact.pk)
                result.added += 1
            elif not renamed:
                result.unchanged += 1

        if to_update:
            Contact.objects.bulk_update(to_update, ["display_name", "attributes", "updated_at"])
        for chunk in _chunks(to_create):
            created = Contact.objects.bulk_create(chunk)
            to_add.extend(contact.pk for contact in created)
            result.created += len(created)
            result.added += len(created)
        for chunk in _chunks(to_add, 1000):
            audience.contacts.add(*chunk)
    return result
