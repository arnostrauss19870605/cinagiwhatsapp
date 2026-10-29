"""Contacts, their full history, and audience management."""

import csv
import io
import json

from django.contrib import messages as flash
from django.contrib.auth.decorators import login_required
from django.db.models import Count, Q
from django.http import HttpResponse
from django.shortcuts import redirect, render
from django.utils import timezone
from django.utils.text import slugify
from django.views.decorators.http import require_POST

from apps.core.audit import audit
from apps.core.scoping import require_role, scoped_get_or_404
from apps.workspaces.models import WorkspaceMembership

from .models import Audience, Contact, normalise_msisdn


def _search_contacts(request, query):
    found = Contact.objects.for_request(request).prefetch_related("audiences")
    if query:
        found = found.filter(
            Q(display_name__icontains=query)
            | Q(profile_name__icontains=query)
            | Q(wa_id__icontains=normalise_msisdn(query) or query)
        )
    return found


@login_required
def contacts(request):
    """Everyone we have ever spoken to, searchable by name or number."""
    if request.workspace is None:
        return redirect("workspaces:list")
    query = (request.GET.get("q") or "").strip()
    found = _search_contacts(request, query)
    return render(
        request,
        "contacts/contacts.html",
        {"contacts": found[:200], "query": query,
         "total": Contact.objects.for_request(request).count()},
    )


def _contacts_xlsx(contacts, *, sheet_title, with_audiences=False):
    """One row per person -> (bytes, row count)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    def local(moment):
        return timezone.localtime(moment).strftime("%Y-%m-%d %H:%M") if moment else ""

    header = ["Name", "Number", "WhatsApp profile name", "Opted out", "Blocked", "First seen", "Last message"]
    widths = [28, 18, 26, 18, 9, 18, 18]
    if with_audiences:
        header.append("Audiences")
        widths.append(40)
    book = Workbook()
    sheet = book.active
    sheet.title = sheet_title
    sheet.append(header)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0065A9")
    count = 0
    for contact in contacts:
        row = [
            contact.name,
            contact.pretty_number,
            contact.profile_name,
            local(contact.opted_out_at),
            "Yes" if contact.is_blocked else "",
            local(contact.first_seen_at),
            local(contact.last_seen_at),
        ]
        if with_audiences:
            row.append(", ".join(a.name for a in contact.audiences.all()))
        sheet.append(row)
        count += 1
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(header))}{max(sheet.max_row, 1)}"
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue(), count


def _xlsx_download(content, filename):
    response = HttpResponse(
        content, content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


@login_required
def contacts_export(request):
    """The contact list as Excel - the whole list, or just what the search
    matched. Supervisors and up, and audited, like a chat export."""
    if request.workspace is None:
        return redirect("workspaces:list")
    require_role(request, *WorkspaceMembership.SUPERVISE_ROLES)
    query = (request.GET.get("q") or "").strip()
    found = _search_contacts(request, query).order_by("display_name", "wa_id")
    content, count = _contacts_xlsx(found, sheet_title="Contacts", with_audiences=True)
    audit("contacts.exported", request=request, query=query, rows=count)
    return _xlsx_download(content, f"contacts-{timezone.localdate():%Y-%m-%d}.xlsx")


def _history(request, contact):
    from apps.inbox.models import Message

    return (
        Message.objects.for_request(request)
        .filter(conversation__contact=contact)
        .select_related("author", "template")
        .order_by("created_at")
    )


@login_required
def contact_detail(request, pk):
    if request.workspace is None:
        return redirect("workspaces:list")
    contact = scoped_get_or_404(Contact, request, pk=pk)
    return render(
        request,
        "contacts/contact_detail.html",
        {
            "contact": contact,
            "history": _history(request, contact),
            "audiences": Audience.objects.for_request(request),
            "member_of": set(contact.audiences.values_list("pk", flat=True)),
            "can_export": request.membership.can_supervise,
            "conversations": contact.conversations.order_by("-last_activity_at"),
        },
    )


@login_required
@require_POST
def contact_audiences(request, pk):
    """Set which audiences this person belongs to - from the chat or their page."""
    contact = scoped_get_or_404(Contact, request, pk=pk)
    chosen = Audience.objects.for_request(request).filter(
        pk__in=request.POST.getlist("audiences")
    )
    contact.audiences.set(chosen)
    audit(
        "contact.audiences_set", request=request, target=contact,
        audiences=[a.name for a in chosen],
    )
    flash.success(request, f"{contact.name} is now in {len(chosen)} audience{'s' if len(chosen) != 1 else ''}.")
    next_url = request.POST.get("next", "")
    if next_url.startswith("/"):
        return redirect(next_url)
    return redirect("contacts:contact_detail", pk=contact.pk)


def _export_rows(history):
    for message in history:
        yield {
            "at": timezone.localtime(message.created_at).strftime("%Y-%m-%d %H:%M:%S"),
            "direction": "in" if message.is_inbound else "out",
            "from": "customer" if message.is_inbound else (
                message.author.display_name if message.author_id else message.get_actor_display()
            ),
            "kind": "internal note" if message.payload.get("note") else message.kind,
            "status": "" if message.is_inbound else message.wa_status,
            "body": message.body or (message.media_filename and f"[{message.media_filename}]") or "",
        }


@login_required
def contact_export(request, pk, fmt):
    """The whole conversation history as a file. Supervisors and up only -
    exporting a customer's chat is data leaving the platform, so it is
    gated and every export is audited."""
    contact = scoped_get_or_404(Contact, request, pk=pk)
    require_role(request, *WorkspaceMembership.SUPERVISE_ROLES)

    history = list(_history(request, contact))
    stem = slugify(contact.name) or contact.wa_id
    audit("contact.exported", request=request, target=contact, fmt=fmt, messages=len(history))

    if fmt == "json":
        payload = {
            "contact": {"name": contact.name, "number": contact.pretty_number},
            "exported_at": timezone.localtime(timezone.now()).isoformat(),
            "messages": list(_export_rows(history)),
        }
        response = HttpResponse(
            json.dumps(payload, indent=2, ensure_ascii=False),
            content_type="application/json; charset=utf-8",
        )
        response["Content-Disposition"] = f'attachment; filename="{stem}-chat.json"'
        return response

    if fmt == "csv":
        response = HttpResponse(content_type="text/csv; charset=utf-8")
        response["Content-Disposition"] = f'attachment; filename="{stem}-chat.csv"'
        writer = csv.DictWriter(
            response, fieldnames=["at", "direction", "from", "kind", "status", "body"]
        )
        writer.writeheader()
        for row in _export_rows(history):
            writer.writerow(row)
        return response

    if fmt == "pdf":
        from .exports import chat_pdf

        response = HttpResponse(content_type="application/pdf")
        response["Content-Disposition"] = f'attachment; filename="{stem}-chat.pdf"'
        chat_pdf(response, request.workspace, contact, history, exported_by=request.user)
        return response

    flash.error(request, "That export format is not available.")
    return redirect("contacts:contact_detail", pk=contact.pk)


def _manager(request):
    require_role(request, *WorkspaceMembership.MANAGE_ROLES)


@login_required
def audiences(request):
    if request.workspace is None:
        return redirect("workspaces:list")
    _manager(request)
    if request.method == "POST":
        name = (request.POST.get("name") or "").strip()
        if not name:
            flash.error(request, "Give the audience a name.")
        elif Audience.objects.for_request(request).filter(name__iexact=name).exists():
            flash.error(request, f"There is already an audience called '{name}'.")
        else:
            audience = Audience.objects.create(
                workspace=request.workspace,
                name=name,
                description=(request.POST.get("description") or "").strip(),
            )
            audit("audience.created", request=request, target=audience, name=name)
            flash.success(request, f"'{name}' created. Now add people to it.")
            return redirect("contacts:audience_detail", pk=audience.pk)
    return render(
        request,
        "contacts/audiences.html",
        {
            "audiences": Audience.objects.for_request(request).annotate(
                member_count=Count("contacts")
            ),
        },
    )


@login_required
def audience_detail(request, pk):
    if request.workspace is None:
        return redirect("workspaces:list")
    _manager(request)
    audience = scoped_get_or_404(Audience, request, pk=pk)

    query = (request.GET.get("q") or "").strip()
    candidates = []
    if query:
        candidates = (
            Contact.objects.for_request(request)
            .filter(Q(display_name__icontains=query) | Q(profile_name__icontains=query)
                    | Q(wa_id__icontains=normalise_msisdn(query) or query))
            .exclude(audiences=audience)[:20]
        )

    return render(
        request,
        "contacts/audience_detail.html",
        {
            "audience": audience,
            "members": audience.contacts.order_by("display_name", "wa_id"),
            "candidates": candidates,
            "query": query,
        },
    )


@login_required
def audience_export(request, pk):
    """Everyone in this audience as an Excel file. Audited, because it is a
    list of customer numbers leaving the platform."""
    _manager(request)
    audience = scoped_get_or_404(Audience, request, pk=pk)
    members = audience.contacts.order_by("display_name", "wa_id")
    content, count = _contacts_xlsx(members, sheet_title="Members")
    audit("audience.exported", request=request, target=audience, rows=count)
    return _xlsx_download(content, f"{slugify(audience.name) or 'audience'}-members.xlsx")


@login_required
@require_POST
def audience_add(request, pk):
    _manager(request)
    audience = scoped_get_or_404(Audience, request, pk=pk)
    contact = scoped_get_or_404(Contact, request, pk=request.POST.get("contact"))
    audience.contacts.add(contact)
    flash.success(request, f"{contact.name} added to '{audience.name}'.")
    return redirect("contacts:audience_detail", pk=audience.pk)


@login_required
@require_POST
def audience_add_new(request, pk):
    """Type in one person who has never messaged us, and put them in this audience.

    Matched on the number, so typing someone who already exists just adds
    them rather than creating a twin.
    """
    _manager(request)
    audience = scoped_get_or_404(Audience, request, pk=pk)
    name = (request.POST.get("name") or "").strip()
    msisdn = normalise_msisdn(request.POST.get("mobile") or "")
    if len(msisdn) < 10:
        flash.error(request, "Enter a mobile number, e.g. 082 123 4567 or +27 82 123 4567.")
        return redirect("contacts:audience_detail", pk=audience.pk)
    contact, created = Contact.objects.get_or_create(
        workspace=request.workspace, wa_id=msisdn, defaults={"display_name": name}
    )
    if not created and name and not contact.display_name:
        contact.display_name = name
        contact.save(update_fields=["display_name"])
    already = audience.contacts.filter(pk=contact.pk).exists()
    audience.contacts.add(contact)
    audit("audience.member_added", request=request, target=audience, contact=contact.pk, created=created)
    if already:
        flash.info(request, f"{contact.name} was already in '{audience.name}'.")
    else:
        flash.success(
            request,
            f"{contact.name} {'added as a new contact and ' if created else ''}put in '{audience.name}'.",
        )
    return redirect("contacts:audience_detail", pk=audience.pk)


@login_required
@require_POST
def audience_remove(request, pk):
    _manager(request)
    audience = scoped_get_or_404(Audience, request, pk=pk)
    contact = scoped_get_or_404(Contact, request, pk=request.POST.get("contact"))
    audience.contacts.remove(contact)
    flash.success(request, f"{contact.name} removed from '{audience.name}'.")
    return redirect("contacts:audience_detail", pk=audience.pk)


@login_required
@require_POST
def audience_delete(request, pk):
    _manager(request)
    audience = scoped_get_or_404(Audience, request, pk=pk)
    name = audience.name
    audience.delete()
    audit("audience.deleted", request=request, name=name)
    flash.success(request, f"'{name}' deleted. The contacts themselves are untouched.")
    return redirect("contacts:audiences")


# --- Importing people from a spreadsheet ------------------------------------
# Step one reads the file and remembers its rows in the session; step two
# shows the columns so the user can say which is which, then imports. Nothing
# is written to the database until the user has seen the mapping.

_IMPORT_KEY = "audience_import"


def _pending_import(request):
    pending = request.session.get(_IMPORT_KEY)
    if not pending or pending.get("workspace") != request.workspace.pk:
        return None
    return pending


@login_required
def audience_import(request):
    """Step one: choose the file and which audience it goes into."""
    from .imports import MAX_ROWS, SheetError, read_sheet

    if request.workspace is None:
        return redirect("workspaces:list")
    _manager(request)
    audiences_here = Audience.objects.for_request(request).annotate(member_count=Count("contacts"))
    form = {
        "target": "existing" if audiences_here.exists() else "new",
        "audience": request.GET.get("audience", ""),
        "name": "",
        "description": "",
        "has_headings": True,
    }
    if request.method == "POST":
        form = {
            "target": "new" if request.POST.get("target") == "new" else "existing",
            "audience": request.POST.get("audience", ""),
            "name": (request.POST.get("name") or "").strip()[:120],
            "description": (request.POST.get("description") or "").strip()[:255],
            "has_headings": request.POST.get("has_headings") == "on",
        }
        error = None
        audience = None
        if form["target"] == "existing":
            audience = audiences_here.filter(pk=form["audience"]).first() if form["audience"].isdigit() else None
            if audience is None:
                error = "Choose the audience these people should go into."
        elif not form["name"]:
            error = "Give the new audience a name."
        elif audiences_here.filter(name__iexact=form["name"]).exists():
            error = f"There is already an audience called '{form['name']}'. Choose it from the list instead."

        upload = request.FILES.get("file")
        if error is None and upload is None:
            error = "Choose a spreadsheet to upload."
        if error is None:
            try:
                headings, rows = read_sheet(upload, has_headings=form["has_headings"])
            except SheetError as exc:
                error = str(exc)
        if error:
            flash.error(request, error)
        else:
            request.session[_IMPORT_KEY] = {
                "workspace": request.workspace.pk,
                "audience": audience.pk if audience else None,
                "new_name": form["name"] if audience is None else "",
                "new_description": form["description"] if audience is None else "",
                "filename": upload.name[:200],
                "headings": headings,
                "rows": rows,
                "first_row": 2 if form["has_headings"] else 1,
            }
            return redirect("contacts:audience_import_map")
    return render(
        request,
        "contacts/audience_import.html",
        {"audiences": audiences_here, "form": form, "max_rows": MAX_ROWS},
    )


@login_required
def audience_import_map(request):
    """Step two: say which column is the number and which are the names, then import."""
    from .imports import FIELDS, guess_mapping, import_rows

    if request.workspace is None:
        return redirect("workspaces:list")
    _manager(request)
    pending = _pending_import(request)
    if pending is None:
        flash.info(request, "Upload the spreadsheet again - that import has already finished or has expired.")
        return redirect("contacts:audience_import")

    audience = None
    if pending["audience"]:
        audience = Audience.objects.for_request(request).filter(pk=pending["audience"]).first()
        if audience is None:
            request.session.pop(_IMPORT_KEY, None)
            flash.error(request, "That audience no longer exists. Start the import again.")
            return redirect("contacts:audience_import")

    headings = pending["headings"]
    columns = list(enumerate(headings))
    mapping = guess_mapping(headings)

    if request.method == "POST":
        def chosen(key):
            raw = request.POST.get(key, "")
            return int(raw) if raw.isdigit() and int(raw) < len(headings) else None

        mapping = {key: chosen(key) for key, _label in FIELDS}
        if mapping["mobile"] is None:
            flash.error(request, "Choose the column that holds the mobile numbers.")
        else:
            if audience is None:
                name = pending["new_name"]
                if Audience.objects.for_request(request).filter(name__iexact=name).exists():
                    request.session.pop(_IMPORT_KEY, None)
                    flash.error(request, f"Someone else has just created an audience called '{name}'. Start the import again.")
                    return redirect("contacts:audience_import")
                audience = Audience.objects.create(
                    workspace=request.workspace, name=name, description=pending["new_description"],
                )
                audit("audience.created", request=request, target=audience, name=name)
            result = import_rows(
                request.workspace, audience, pending["rows"], mapping,
                first_row_number=pending["first_row"],
            )
            request.session.pop(_IMPORT_KEY, None)
            audit(
                "audience.imported", request=request, target=audience,
                file=pending["filename"],
                columns={key: headings[index] for key, index in mapping.items() if index is not None},
                added=result.added, created=result.created, renamed=len(result.updated),
                unchanged=result.unchanged, repeated=result.repeated, unusable=len(result.unusable),
            )
            return render(
                request,
                "contacts/audience_import_done.html",
                {"audience": audience, "result": result, "filename": pending["filename"]},
            )

    fields = [
        {"key": key, "label": label, "selected": mapping.get(key)}
        for key, label in FIELDS
    ]
    return render(
        request,
        "contacts/audience_import_map.html",
        {
            "pending": pending,
            "audience": audience,
            "columns": columns,
            "fields": fields,
            "preview": pending["rows"][:5],
            "row_count": len(pending["rows"]),
        },
    )
