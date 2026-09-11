"""Questions asked over WhatsApp, their answers, and the prize draw."""

import io
import secrets
from collections import Counter

from django.contrib import messages as flash
from django.contrib.auth.decorators import login_required
from django.core.paginator import Paginator
from django.db import transaction
from django.db.models import Count, Sum
from django.http import HttpResponse
from django.shortcuts import redirect, render
from django.utils import timezone
from django.utils.text import slugify
from django.views.decorators.http import require_POST

from apps.contacts.models import Audience, Contact
from apps.core.audit import audit
from apps.core.scoping import require_role, scoped_get_or_404
from apps.library.models import MessageTemplate
from apps.workspaces.models import WorkspaceMembership

from .models import Answer, AnswerOption, BonusEntry, PrizeDraw, Question

PAGE_SIZE = 50


def _manager(request):
    require_role(request, *WorkspaceMembership.MANAGE_ROLES)


def _supervisor(request):
    require_role(request, *WorkspaceMembership.SUPERVISE_ROLES)


@login_required
def questions(request):
    if request.workspace is None:
        return redirect("workspaces:list")
    _supervisor(request)
    rows = Question.objects.for_request(request).select_related("template").annotate(answered=Count("answers"))
    return render(request, "questions/questions.html", {"questions": rows})


def _candidate_templates(request, current=None):
    """Approved templates with quick-reply buttons that are not already a question."""
    taken = Question.objects.for_request(request).exclude(pk=getattr(current, "pk", None)).values_list("template_id", flat=True)
    return [
        t for t in MessageTemplate.objects.for_request(request)
        .filter(status=MessageTemplate.Status.APPROVED).exclude(pk__in=taken).order_by("name")
        if t.quick_reply_labels
    ]


@login_required
def question_edit(request, pk=None):
    if request.workspace is None:
        return redirect("workspaces:list")
    _manager(request)
    question = scoped_get_or_404(Question, request, pk=pk) if pk else None
    templates = _candidate_templates(request, question)
    audiences = Audience.objects.for_request(request)

    chosen = question.template if question else None
    template_id = request.POST.get("template") or request.GET.get("template")
    if template_id and (question is None):
        chosen = next((t for t in templates if str(t.pk) == str(template_id)), None)

    labels = chosen.quick_reply_labels if chosen else []
    existing = {o.label: o for o in question.options.all()} if question else {}
    option_rows = []
    for index, label in enumerate(labels):
        current = existing.get(label)
        prefix = f"opt_{index}"
        option_rows.append({
            "label": label,
            "prefix": prefix,
            "value": request.POST.get(f"{prefix}_value", current.value if current else slugify(label).replace("-", "_")[:60]),
            "score": request.POST.get(f"{prefix}_score", current.score if current else 0),
            "is_correct": request.POST.get(f"{prefix}_correct") == "on" if request.method == "POST" else bool(current and current.is_correct),
            "audience": request.POST.get(f"{prefix}_audience", str(current.audience_id) if current and current.audience_id else ""),
        })

    error = None
    if request.method == "POST" and request.POST.get("save"):
        title = (request.POST.get("title") or "").strip()
        thanks = (request.POST.get("thanks_text") or "").strip()
        counts = request.POST.get("counts_as_entry") == "on"
        if chosen is None:
            error = "Choose an approved message with reply buttons."
        elif not title:
            error = "Give the question a name for the team."
        else:
            values = [slugify(r["value"]).replace("-", "_")[:60] or f"option_{i}" for i, r in enumerate(option_rows)]
            if len(set(values)) != len(values):
                error = "Each answer needs a different short code."
        if error is None:
            with transaction.atomic():
                if question is None:
                    question = Question(workspace=request.workspace, template=chosen, created_by=request.user)
                question.title = title
                question.thanks_text = thanks
                question.counts_as_entry = counts
                question.save()
                seen = []
                for index, row in enumerate(option_rows):
                    try:
                        score = int(row["score"] or 0)
                    except (TypeError, ValueError):
                        score = 0
                    audience = audiences.filter(pk=row["audience"]).first() if row["audience"] else None
                    option, _ = AnswerOption.objects.update_or_create(
                        question=question, label=row["label"],
                        defaults={
                            "value": values[index], "score": score,
                            "is_correct": row["is_correct"], "audience": audience, "order": index,
                        },
                    )
                    seen.append(option.pk)
                question.options.exclude(pk__in=seen).delete()
            audit("question.saved", request=request, target=question, template=chosen.name)
            flash.success(request, f"'{question.title}' is ready. Send '{chosen.name}' to an audience and answers will be counted here.")
            return redirect("questions:question_results", pk=question.pk)
        flash.error(request, error)

    return render(
        request,
        "questions/question_form.html",
        {
            "question": question,
            "templates": templates,
            "chosen": chosen,
            "option_rows": option_rows,
            "audiences": audiences,
            "title": request.POST.get("title", question.title if question else ""),
            "thanks_text": request.POST.get("thanks_text", question.thanks_text if question else ""),
            "counts_as_entry": (request.POST.get("counts_as_entry") == "on") if request.method == "POST" else (question.counts_as_entry if question else True),
        },
    )


@login_required
def question_results(request, pk):
    if request.workspace is None:
        return redirect("workspaces:list")
    _supervisor(request)
    question = scoped_get_or_404(Question, request, pk=pk)
    options = list(question.options.annotate(n=Count("answers")))
    answered = question.people_answered
    for option in options:
        option.pct = round(option.n * 100 / answered) if answered else 0
    answers = question.answers.select_related("contact", "option", "bulk_send").order_by("-answered_at")
    page = Paginator(answers, PAGE_SIZE).get_page(request.GET.get("page"))
    return render(
        request,
        "questions/question_results.html",
        {
            "question": question,
            "options": options,
            "asked": question.people_asked,
            "answered": answered,
            "unmatched": question.answers.filter(option__isnull=True).count(),
            "correct": question.answers.filter(is_correct=True).count() if any(o.is_correct for o in options) else None,
            "page": page,
            "answers": page.object_list,
        },
    )


@login_required
@require_POST
def question_toggle(request, pk):
    _manager(request)
    question = scoped_get_or_404(Question, request, pk=pk)
    question.is_active = not question.is_active
    question.save(update_fields=["is_active"])
    flash.success(request, f"'{question.title}' is {'counting answers again' if question.is_active else 'switched off'}.")
    return redirect("questions:question_results", pk=question.pk)


def _xlsx(sheets):
    """sheets: [(title, header_row, rows)] -> bytes."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    book = Workbook()
    first = True
    for title, header, rows in sheets:
        sheet = book.active if first else book.create_sheet()
        first = False
        sheet.title = title
        sheet.append(header)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="0065A9")
        for row in rows:
            sheet.append(row)
        sheet.freeze_panes = "A2"
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _download(content, filename):
    response = HttpResponse(
        content, content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


@login_required
def question_export(request, pk):
    _supervisor(request)
    question = scoped_get_or_404(Question, request, pk=pk)
    answers = question.answers.select_related("contact", "option").order_by("-answered_at")
    rows = [
        [
            a.contact.name, a.contact.pretty_number, a.option.label if a.option else a.raw_text,
            a.option.value if a.option else "", a.score, "Yes" if a.is_correct else "",
            timezone.localtime(a.first_answered_at).strftime("%Y-%m-%d %H:%M"), a.taps,
        ]
        for a in answers
    ]
    audit("question.exported", request=request, target=question, rows=len(rows))
    return _download(
        _xlsx([(
            "Answers",
            ["Name", "Number", "Answer", "Code", "Score", "Correct", "Answered", "Taps"],
            rows,
        )]),
        f"{slugify(question.title) or 'question'}-answers.xlsx",
    )


# -- prize draw ---------------------------------------------------------------


def _entries(request):
    """Entries per contact: one per answered question that counts, plus bonuses."""
    counts = Counter()
    answered = (
        Answer.objects.for_request(request)
        .filter(question__counts_as_entry=True, question__is_active=True)
        .values_list("contact_id", flat=True)
    )
    for contact_id in answered:
        counts[contact_id] += 1
    bonuses = (
        BonusEntry.objects.for_request(request).values("contact_id").annotate(n=Sum("entries"))
    )
    for row in bonuses:
        counts[row["contact_id"]] += row["n"]
    contacts = {c.pk: c for c in Contact.objects.for_request(request).filter(pk__in=counts)}
    rows = [
        {"contact": contacts[pk], "entries": n}
        for pk, n in counts.items() if pk in contacts and not contacts[pk].is_opted_out
    ]
    rows.sort(key=lambda r: (-r["entries"], r["contact"].name))
    return rows


@login_required
def prize_draw(request):
    if request.workspace is None:
        return redirect("workspaces:list")
    _supervisor(request)
    rows = _entries(request)
    return render(
        request,
        "questions/prize_draw.html",
        {
            "rows": rows,
            "total_entries": sum(r["entries"] for r in rows),
            "questions": Question.objects.for_request(request).filter(counts_as_entry=True, is_active=True),
            "draws": PrizeDraw.objects.for_request(request).select_related("winner", "drawn_by"),
            "bonuses": BonusEntry.objects.for_request(request).select_related("contact", "created_by")[:20],
        },
    )


@login_required
@require_POST
def bonus_entry(request):
    _manager(request)
    contact = scoped_get_or_404(Contact, request, pk=request.POST.get("contact"))
    try:
        entries = max(1, min(10, int(request.POST.get("entries") or 1)))
    except ValueError:
        entries = 1
    reason = (request.POST.get("reason") or "Bonus entry").strip()[:120]
    BonusEntry.objects.create(
        workspace=request.workspace, contact=contact, entries=entries, reason=reason, created_by=request.user
    )
    audit("prize_draw.bonus", request=request, target=contact, entries=entries, reason=reason)
    flash.success(request, f"{contact.name} has {entries} extra entr{'y' if entries == 1 else 'ies'}.")
    return redirect("questions:prize_draw")


@login_required
@require_POST
def draw_winner(request):
    """Pick one entry at random and record it. Every entry has an equal chance."""
    _manager(request)
    rows = _entries(request)
    if not rows:
        flash.error(request, "Nobody has any entries yet.")
        return redirect("questions:prize_draw")
    name = (request.POST.get("name") or "Prize draw").strip()[:120]
    pool = [row for row in rows for _ in range(row["entries"])]
    winner = secrets.choice(pool)
    draw = PrizeDraw.objects.create(
        workspace=request.workspace, name=name, winner=winner["contact"],
        winner_entries=winner["entries"], total_entries=len(pool), people=len(rows), drawn_by=request.user,
    )
    audit("prize_draw.drawn", request=request, target=draw, winner=winner["contact"].pk, total=len(pool))
    flash.success(request, f"{winner['contact'].name} wins '{name}', drawn from {len(pool)} entries across {len(rows)} people.")
    return redirect("questions:prize_draw")


@login_required
def prize_draw_export(request):
    _supervisor(request)
    rows = _entries(request)
    audit("prize_draw.exported", request=request, rows=len(rows))
    return _download(
        _xlsx([(
            "Entries",
            ["Name", "Number", "Entries"],
            [[r["contact"].name, r["contact"].pretty_number, r["entries"]] for r in rows],
        )]),
        "prize-draw-entries.xlsx",
    )
