"""Analytics: the overview figures and the bulk-send drill-down."""

import datetime as dt
from unittest import mock

from django.contrib.auth import get_user_model
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone

from apps.channels_wa.messaging.base import SendResult
from apps.channels_wa.models import WhatsAppChannel
from apps.contacts.models import Audience, Contact
from apps.inbox.models import Conversation, Message
from apps.library.models import BulkSend, MessageTemplate
from apps.library.tasks import run_bulk_send
from apps.workspaces.models import Workspace, WorkspaceMembership

User = get_user_model()
SENT = SendResult(ok=True, wamid="wamid.BULK")


@override_settings(OUTBOUND_COMMS_MODE="live", CELERY_TASK_ALWAYS_EAGER=True)
class ReportingTests(TestCase):
    def setUp(self):
        self.workspace = Workspace.objects.create(name="Cinagi Broker Support")
        self.manager = User.objects.create_user("boss", "boss@example.com", "pw12345678")
        WorkspaceMembership.objects.create(user=self.manager, workspace=self.workspace, role="owner")
        self.agent = User.objects.create_user("agent", "agent@example.com", "pw12345678")
        WorkspaceMembership.objects.create(user=self.agent, workspace=self.workspace, role="agent")
        self.channel = WhatsAppChannel.objects.create(
            workspace=self.workspace, display_name="Broker Support", phone_number_id="123", is_active=True
        )
        self.template = MessageTemplate.objects.create(
            workspace=self.workspace,
            channel=self.channel,
            name="event_invite",
            language="en",
            status=MessageTemplate.Status.APPROVED,
            components=[{"type": "BODY", "text": "Hi {{1}}, join us on {{2}}."}],
        )
        self.thabo = Contact.objects.create(
            workspace=self.workspace, wa_id="27726124698", display_name="Thabo Mokoena"
        )
        self.ann = Contact.objects.create(
            workspace=self.workspace, wa_id="27831234567", display_name="Ann Botha"
        )
        self.audience = Audience.objects.create(workspace=self.workspace, name="Launch attendees")
        self.audience.contacts.set([self.thabo, self.ann])
        self.client.force_login(self.manager)

    def _run_bulk_send(self):
        with mock.patch(
            "apps.channels_wa.messaging.meta_cloud.MetaCloudChannel.send_template", return_value=SENT
        ):
            run_bulk_send(
                self.workspace.pk, self.template.pk, [self.audience.pk], ["Wednesday"], None,
                self.manager.pk,
            )
        return BulkSend.objects.get()

    def test_a_bulk_send_is_recorded_and_its_messages_point_back_to_it(self):
        batch = self._run_bulk_send()
        self.assertEqual(batch.status, BulkSend.Status.DONE)
        self.assertEqual(batch.template_name, "event_invite")
        self.assertEqual(batch.audience_names, ["Launch attendees"])
        self.assertEqual(batch.recipient_count, 2)
        self.assertEqual(batch.sent_count, 2)
        self.assertEqual(batch.created_by, self.manager)
        self.assertEqual(batch.messages.count(), 2)

    def test_the_drill_down_reflects_receipts_and_replies(self):
        batch = self._run_bulk_send()
        thabo_message = batch.messages.get(conversation__contact=self.thabo)
        thabo_message.advance_status("read")
        thabo_message.save()
        Message.objects.create(
            workspace=self.workspace,
            conversation=thabo_message.conversation,
            direction=Message.Direction.IN,
            body="Count me in",
        )
        ann_message = batch.messages.get(conversation__contact=self.ann)
        ann_message.wa_status = Message.Status.FAILED
        ann_message.wa_error = {"code": 131026, "title": "Message undeliverable"}
        ann_message.save()

        response = self.client.get(reverse("reporting:bulk_send", args=[batch.pk]))
        self.assertEqual(response.status_code, 200)
        stats = response.context["stats"]
        self.assertEqual((stats["sent"], stats["read"], stats["replied"], stats["failed"]), (1, 1, 1, 1))
        self.assertEqual(stats["read_pct"], 100)
        self.assertContains(response, "Thabo Mokoena")
        self.assertContains(response, "cannot receive WhatsApp messages")

        failed_only = self.client.get(reverse("reporting:bulk_send", args=[batch.pk]) + "?show=failed")
        self.assertEqual([m.conversation.contact for m in failed_only.context["recipients"]], [self.ann])

    def test_the_overview_counts_sends_replies_and_consultant_requests(self):
        batch = self._run_bulk_send()
        message = batch.messages.first()
        message.advance_status("delivered")
        message.save()
        Message.objects.create(
            workspace=self.workspace,
            conversation=message.conversation,
            direction=Message.Direction.IN,
            body="Cinagi Consultant",
        )
        Message.objects.create(
            workspace=self.workspace,
            conversation=message.conversation,
            direction=Message.Direction.SYSTEM,
            actor=Message.Actor.SYSTEM,
            body="Customer asked to speak to a person.",
            payload={"note": True, "escalation": "human_requested"},
        )
        Message.objects.create(
            workspace=self.workspace,
            conversation=message.conversation,
            direction=Message.Direction.OUT,
            actor=Message.Actor.AGENT,
            author=self.agent,
            body="Hello, how can I help?",
            wa_status=Message.Status.SENT,
        )

        response = self.client.get(reverse("reporting:overview"))
        self.assertEqual(response.status_code, 200)
        totals = response.context["totals"]
        self.assertEqual(totals["sent"], 3)
        self.assertEqual(totals["delivered"], 1)
        self.assertEqual(totals["inbound"], 1)
        self.assertEqual(totals["consultant_requests"], 1)
        self.assertEqual(totals["new_chats"], 2)
        self.assertEqual(len(response.context["daily"]), 30)
        self.assertEqual(sum(d["out"] for d in response.context["daily"]), 3)
        template_row = response.context["template_rows"][0]
        self.assertEqual((template_row["name"], template_row["sent"], template_row["replied"]), ("event_invite", 2, 1))
        self.assertEqual(response.context["agent_rows"][0]["replies"], 1)
        self.assertEqual(len(response.context["batches"]), 1)
        self.assertContains(response, "Launch attendees")

    def test_the_period_switch_is_bounded(self):
        response = self.client.get(reverse("reporting:overview") + "?days=7")
        self.assertEqual(len(response.context["daily"]), 7)
        response = self.client.get(reverse("reporting:overview") + "?days=999")
        self.assertEqual(response.context["days"], 30)

    def test_agents_cannot_see_analytics(self):
        self.client.force_login(self.agent)
        self.assertEqual(self.client.get(reverse("reporting:overview")).status_code, 403)

    def test_another_workspaces_bulk_send_is_not_found(self):
        batch = self._run_bulk_send()
        other = Workspace.objects.create(name="Other")
        WorkspaceMembership.objects.create(user=self.manager, workspace=other, role="owner")
        self.client.get(reverse("workspaces:switch", args=[other.slug]))
        self.assertEqual(self.client.get(reverse("reporting:bulk_send", args=[batch.pk])).status_code, 404)

    def test_pressing_send_records_the_batch_before_the_background_run(self):
        with mock.patch("apps.library.tasks.run_bulk_send.delay") as delay:
            response = self.client.post(
                reverse("library:bulk_send"),
                {"template": self.template.pk, "audiences": [self.audience.pk], "value_2": "Wednesday"},
            )
        batch = BulkSend.objects.get()
        self.assertEqual(batch.status, BulkSend.Status.QUEUED)
        self.assertEqual(batch.recipient_count, 2)
        self.assertEqual(delay.call_args.args[-1], batch.pk)
        self.assertRedirects(response, reverse("reporting:bulk_send", args=[batch.pk]))


@override_settings(OUTBOUND_COMMS_MODE="live", CELERY_TASK_ALWAYS_EAGER=True)
class BulkSendExportAndResendTests(ReportingTests):
    """Reuses the ReportingTests fixture: an owner, a template, two contacts in one audience."""

    def _batch_with_outcomes(self):
        batch = self._run_bulk_send()
        thabo = batch.messages.get(conversation__contact=self.thabo)
        thabo.advance_status("read")
        thabo.save()
        ann = batch.messages.get(conversation__contact=self.ann)
        ann.wa_status = Message.Status.FAILED
        ann.wa_error = {"code": 131026, "title": "Message undeliverable"}
        ann.save()
        return batch

    def test_excel_export_lists_every_recipient_and_respects_the_filter(self):
        from io import BytesIO

        from openpyxl import load_workbook

        batch = self._batch_with_outcomes()
        response = self.client.get(reverse("reporting:bulk_send_export", args=[batch.pk, "xlsx"]))
        self.assertEqual(response.status_code, 200)
        self.assertIn("spreadsheetml", response["Content-Type"])
        book = load_workbook(BytesIO(response.content))
        rows = list(book["Recipients"].iter_rows(values_only=True))
        self.assertEqual(rows[0][:3], ("Name", "Number", "Status"))
        self.assertEqual(len(rows), 3)
        self.assertEqual(book["Summary"]["A1"].value, "Message")

        filtered = self.client.get(reverse("reporting:bulk_send_export", args=[batch.pk, "xlsx"]) + "?show=failed")
        rows = list(load_workbook(BytesIO(filtered.content))["Recipients"].iter_rows(values_only=True))
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[1][0], "Ann Botha")
        self.assertIn("cannot receive WhatsApp", rows[1][7])

    def test_pdf_report_is_produced(self):
        batch = self._batch_with_outcomes()
        response = self.client.get(reverse("reporting:bulk_send_export", args=[batch.pk, "pdf"]))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/pdf")
        self.assertTrue(response.content.startswith(b"%PDF"))

    def test_the_recipient_list_is_paginated(self):
        batch = self._run_bulk_send()
        with mock.patch("apps.reporting.views.PAGE_SIZE", 1):
            first = self.client.get(reverse("reporting:bulk_send", args=[batch.pk]))
            second = self.client.get(reverse("reporting:bulk_send", args=[batch.pk]) + "?page=2")
        self.assertEqual(len(first.context["recipients"]), 1)
        self.assertEqual(first.context["page"].paginator.num_pages, 2)
        self.assertContains(first, "Page 1 of 2")
        self.assertEqual(len(second.context["recipients"]), 1)
        self.assertNotEqual(first.context["recipients"][0].pk, second.context["recipients"][0].pk)

    def test_resending_to_not_delivered_makes_a_new_batch_for_just_those_people(self):
        batch = self._batch_with_outcomes()
        with mock.patch(
            "apps.channels_wa.messaging.meta_cloud.MetaCloudChannel.send_template", return_value=SENT
        ) as send:
            response = self.client.post(
                reverse("reporting:bulk_send_resend", args=[batch.pk]), {"group": "failed"}
            )
        new = BulkSend.objects.exclude(pk=batch.pk).get()
        self.assertRedirects(response, reverse("reporting:bulk_send", args=[new.pk]))
        self.assertEqual(new.status, BulkSend.Status.DONE)
        self.assertEqual(new.template, self.template)
        self.assertEqual(new.values, batch.values)
        self.assertIn("not delivered", new.audience_names[0])
        self.assertEqual(new.recipient_count, 1)
        send.assert_called_once()
        self.assertEqual(send.call_args.args[0], self.ann.wa_id)
        self.assertEqual(new.messages.get().conversation.contact, self.ann)

    def test_resending_to_not_read_skips_read_and_failed_people(self):
        batch = self._batch_with_outcomes()
        # A third person who was delivered to but never opened it.
        sipho = Contact.objects.create(workspace=self.workspace, wa_id="27840000003", display_name="Sipho")
        conversation = Conversation.objects.create(workspace=self.workspace, channel=self.channel, contact=sipho)
        Message.objects.create(
            workspace=self.workspace, conversation=conversation, direction=Message.Direction.OUT,
            actor=Message.Actor.BOT, kind=Message.Kind.TEMPLATE, template=self.template,
            bulk_send=batch, wa_status=Message.Status.DELIVERED,
        )
        with mock.patch(
            "apps.channels_wa.messaging.meta_cloud.MetaCloudChannel.send_template", return_value=SENT
        ) as send:
            self.client.post(reverse("reporting:bulk_send_resend", args=[batch.pk]), {"group": "unread"})
        send.assert_called_once()
        self.assertEqual(send.call_args.args[0], sipho.wa_id)

    def test_opted_out_people_are_left_out_of_a_resend(self):
        batch = self._batch_with_outcomes()
        self.ann.opted_out_at = timezone.now()
        self.ann.save()
        response = self.client.post(reverse("reporting:bulk_send_resend", args=[batch.pk]), {"group": "failed"})
        self.assertRedirects(response, reverse("reporting:bulk_send", args=[batch.pk]))
        self.assertEqual(BulkSend.objects.count(), 1)

    def test_resend_needs_the_bulk_right(self):
        batch = self._batch_with_outcomes()
        WorkspaceMembership.objects.filter(user=self.agent).update(role="supervisor")
        self.client.force_login(self.agent)
        page = self.client.get(reverse("reporting:bulk_send", args=[batch.pk]))
        self.assertEqual(page.status_code, 200)
        self.assertFalse(page.context["can_resend"])
        response = self.client.post(reverse("reporting:bulk_send_resend", args=[batch.pk]), {"group": "failed"})
        self.assertEqual(response.status_code, 403)

    def test_a_no_longer_approved_template_cannot_be_resent(self):
        batch = self._batch_with_outcomes()
        self.template.status = MessageTemplate.Status.PAUSED
        self.template.save()
        response = self.client.post(reverse("reporting:bulk_send_resend", args=[batch.pk]), {"group": "failed"})
        self.assertRedirects(response, reverse("reporting:bulk_send", args=[batch.pk]))
        self.assertEqual(BulkSend.objects.count(), 1)
