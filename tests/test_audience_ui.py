"""The audience pages and the bulk send screen.

Bulk sending from a browser is the platform's biggest red button, so the
tests pin down who may press it (managers), that pressing it only queues
background work, and that a spreadsheet upload can never duplicate a person.
"""

from io import BytesIO
from unittest import mock

from django.contrib.auth import get_user_model
from django.test import TestCase, override_settings
from django.urls import reverse

from apps.channels_wa.messaging.base import SendResult
from apps.channels_wa.models import WhatsAppChannel
from apps.contacts.models import Audience, Contact
from apps.inbox.models import Message
from apps.library.models import MessageTemplate
from apps.workspaces.models import Workspace, WorkspaceMembership

User = get_user_model()


class AudienceUiBase(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.workspace = Workspace.objects.create(name="Cinagi Broker Support")
        cls.manager = User.objects.create_user("manager", "m@example.com", "pw12345678")
        cls.agent = User.objects.create_user("agent", "a@example.com", "pw12345678")
        WorkspaceMembership.objects.create(
            user=cls.manager, workspace=cls.workspace, role="admin", can_send_bulk=True
        )
        WorkspaceMembership.objects.create(user=cls.agent, workspace=cls.workspace, role="agent")
        cls.channel = WhatsAppChannel.objects.create(
            workspace=cls.workspace, display_name="Broker Support",
            phone_number="+27842358896", phone_number_id="123", is_active=True,
        )

    def setUp(self):
        self.client.force_login(self.manager)


class AudiencePageTests(AudienceUiBase):
    def test_a_manager_can_create_an_audience(self):
        response = self.client.post(
            reverse("contacts:audiences"), {"name": "Launch attendees", "description": "RSVP yes"}
        )
        audience = Audience.objects.get(name="Launch attendees")
        self.assertRedirects(response, reverse("contacts:audience_detail", args=[audience.pk]))

    def test_an_agent_may_not_manage_audiences(self):
        self.client.force_login(self.agent)
        self.assertEqual(self.client.get(reverse("contacts:audiences")).status_code, 403)

    def test_duplicate_names_are_refused_in_plain_language(self):
        Audience.objects.create(workspace=self.workspace, name="Launch attendees")
        response = self.client.post(reverse("contacts:audiences"), {"name": "launch ATTENDEES"})
        self.assertEqual(Audience.objects.count(), 1)
        self.assertContains(response, "already an audience")

    def test_members_can_be_added_and_removed(self):
        audience = Audience.objects.create(workspace=self.workspace, name="VIPs")
        contact = Contact.objects.create(
            workspace=self.workspace, wa_id="27726124698", display_name="Thabo Mokoena"
        )
        self.client.post(
            reverse("contacts:audience_add", args=[audience.pk]), {"contact": contact.pk}
        )
        self.assertEqual(audience.contacts.count(), 1)
        self.client.post(
            reverse("contacts:audience_remove", args=[audience.pk]), {"contact": contact.pk}
        )
        self.assertEqual(audience.contacts.count(), 0)

    def test_members_can_be_downloaded_as_excel(self):
        from openpyxl import load_workbook

        from apps.core.models import AuditLog

        audience = Audience.objects.create(workspace=self.workspace, name="VIP brokers")
        audience.contacts.add(Contact.objects.create(
            workspace=self.workspace, wa_id="27726124698", display_name="Thabo Mokoena"
        ))
        Contact.objects.create(workspace=self.workspace, wa_id="27831234567", display_name="Not a member")
        response = self.client.get(reverse("contacts:audience_export", args=[audience.pk]))
        self.assertEqual(response.status_code, 200)
        self.assertIn('filename="vip-brokers-members.xlsx"', response["Content-Disposition"])
        rows = list(load_workbook(BytesIO(response.content)).active.values)
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[1][:2], ("Thabo Mokoena", "+27726124698"))
        self.assertTrue(AuditLog.objects.filter(action="audience.exported").exists())

    def test_an_agent_may_not_download_an_audience(self):
        audience = Audience.objects.create(workspace=self.workspace, name="VIPs")
        self.client.force_login(self.agent)
        response = self.client.get(reverse("contacts:audience_export", args=[audience.pk]))
        self.assertEqual(response.status_code, 403)

    def test_another_workspaces_audience_cannot_be_downloaded(self):
        other = Workspace.objects.create(name="Bravo")
        audience = Audience.objects.create(workspace=other, name="Theirs")
        response = self.client.get(reverse("contacts:audience_export", args=[audience.pk]))
        self.assertEqual(response.status_code, 404)

    def test_a_contact_from_another_workspace_cannot_be_added(self):
        other = Workspace.objects.create(name="Bravo")
        stranger = Contact.objects.create(workspace=other, wa_id="27820000009")
        audience = Audience.objects.create(workspace=self.workspace, name="VIPs")
        response = self.client.post(
            reverse("contacts:audience_add", args=[audience.pk]), {"contact": stranger.pk}
        )
        self.assertEqual(response.status_code, 404)
        self.assertEqual(audience.contacts.count(), 0)


def _xlsx(rows):
    from openpyxl import Workbook

    book = Workbook()
    for row in rows:
        book.active.append(row)
    buffer = BytesIO()
    book.save(buffer)
    buffer.seek(0)
    buffer.name = "guests.xlsx"
    return buffer


class AudienceImportTests(AudienceUiBase):
    """Spreadsheet import: upload, match columns, import.

    The sheet's columns are deliberately not called "name" and "mobile" -
    the point of the mapping step is that people's own spreadsheets work.
    """

    SHEET = [
        ["Ref", "Naam", "Van", "Cell No"],
        ["A1", "Thabo", "Mokoena", "072 612 4698"],
        ["A2", "Ann", "Botha", 831234567],  # Excel dropped the leading zero
        ["A3", "No", "Number", ""],
        ["A4", "Thabo", "Again", "+27 72 612 4698"],
    ]
    MAPPING = {"mobile": "3", "first_name": "1", "last_name": "2", "full_name": ""}

    def setUp(self):
        super().setUp()
        self.audience = Audience.objects.create(workspace=self.workspace, name="Launch attendees")

    def _upload(self, rows=None, **form):
        data = {"target": "existing", "audience": self.audience.pk, "has_headings": "on",
                "file": _xlsx(rows or self.SHEET)}
        data.update(form)
        return self.client.post(reverse("contacts:audience_import"), data)

    def _import(self, mapping=None, **form):
        self._upload(**form)
        return self.client.post(reverse("contacts:audience_import_map"), mapping or self.MAPPING)

    def test_upload_shows_the_columns_with_a_sensible_guess(self):
        response = self._upload()
        self.assertRedirects(response, reverse("contacts:audience_import_map"))
        page = self.client.get(reverse("contacts:audience_import_map"))
        self.assertContains(page, "Cell No")
        self.assertEqual(
            {f["key"]: f["selected"] for f in page.context["fields"]},
            {"mobile": 3, "full_name": None, "first_name": 1, "last_name": 2},
        )
        self.assertEqual(Contact.objects.count(), 0, "nothing is written before the mapping")

    def test_mapped_columns_create_contacts_and_membership(self):
        response = self._import()
        self.assertEqual(response.status_code, 200)
        result = response.context["result"]
        self.assertEqual((result.added, result.created, result.repeated), (2, 2, 1))
        self.assertEqual([row for row, _ in result.unusable], [4])
        thabo = Contact.objects.get(wa_id="27726124698")
        self.assertEqual(thabo.display_name, "Thabo Mokoena")
        self.assertEqual(thabo.attributes, {"first_name": "Thabo", "last_name": "Mokoena"})
        self.assertTrue(self.audience.contacts.filter(wa_id="27831234567").exists())

    def test_existing_members_are_skipped_or_renamed(self):
        same = Contact.objects.create(workspace=self.workspace, wa_id="27726124698", display_name="Thabo Mokoena")
        renamed = Contact.objects.create(workspace=self.workspace, wa_id="27831234567", display_name="Ann Smit")
        self.audience.contacts.add(same, renamed)
        result = self._import().context["result"]
        self.assertEqual((result.added, result.unchanged), (0, 1))
        self.assertEqual(result.updated, [("+27831234567", "Ann Smit", "Ann Botha")])
        renamed.refresh_from_db()
        self.assertEqual(renamed.display_name, "Ann Botha")
        self.assertEqual(Contact.objects.count(), 2)
        self.assertEqual(self.audience.contacts.count(), 2)

    def test_a_contact_outside_the_audience_is_added_not_duplicated(self):
        Contact.objects.create(workspace=self.workspace, wa_id="27726124698", profile_name="T")
        result = self._import().context["result"]
        self.assertEqual((result.added, result.created), (2, 1))
        self.assertEqual(Contact.objects.filter(wa_id="27726124698").count(), 1)

    def test_a_new_audience_is_only_created_when_the_import_runs(self):
        self._upload(target="new", name="Press list", audience="")
        self.assertFalse(Audience.objects.filter(name="Press list").exists())
        self.client.post(reverse("contacts:audience_import_map"), self.MAPPING)
        self.assertEqual(Audience.objects.get(name="Press list").contacts.count(), 2)

    def test_a_new_audience_may_not_reuse_a_name(self):
        self._upload(target="new", name="launch ATTENDEES", audience="")
        response = self.client.get(reverse("contacts:audience_import_map"))
        self.assertRedirects(response, reverse("contacts:audience_import"))

    def test_a_single_name_column_is_read_as_the_whole_name(self):
        from apps.contacts.imports import guess_mapping

        self.assertEqual(
            guess_mapping(["Name", "Cell"]),
            {"mobile": 1, "full_name": 0, "first_name": None, "last_name": None},
        )

    def test_the_number_column_is_required(self):
        self._upload()
        self.client.post(reverse("contacts:audience_import_map"), {**self.MAPPING, "mobile": ""})
        self.assertEqual(Contact.objects.count(), 0)

    def test_a_csv_still_works(self):
        upload = BytesIO(b"name;cell\nThabo Mokoena;0726124698\n")
        upload.name = "guests.csv"
        self._upload(file=upload)
        self.client.post(reverse("contacts:audience_import_map"), {"mobile": "1", "full_name": "0"})
        self.assertEqual(self.audience.contacts.get().display_name, "Thabo Mokoena")

    def test_an_old_excel_file_is_refused_in_plain_language(self):
        upload = BytesIO(b"junk")
        upload.name = "guests.xls"
        response = self._upload(file=upload)
        self.assertContains(response, "old-style Excel file")

    def test_an_agent_may_not_import(self):
        self.client.force_login(self.agent)
        self.assertEqual(self._upload().status_code, 403)

    def test_another_workspaces_audience_cannot_be_the_target(self):
        other = Workspace.objects.create(name="Elsewhere")
        theirs = Audience.objects.create(workspace=other, name="Theirs")
        response = self._upload(audience=theirs.pk)
        self.assertEqual(response.status_code, 200)
        self.assertNotIn("audience_import", self.client.session)

    def test_the_mapping_step_without_an_upload_starts_over(self):
        response = self.client.get(reverse("contacts:audience_import_map"))
        self.assertRedirects(response, reverse("contacts:audience_import"))


@override_settings(OUTBOUND_COMMS_MODE="live", CELERY_TASK_ALWAYS_EAGER=True)
class BulkSendUiTests(AudienceUiBase):
    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        cls.template = MessageTemplate.objects.create(
            workspace=cls.workspace, channel=cls.channel, name="logistics_update",
            language="en", status=MessageTemplate.Status.APPROVED,
            components=[{"type": "BODY", "text": "Hi {{1}}, we are on for tomorrow."}],
        )
        cls.audience = Audience.objects.create(workspace=cls.workspace, name="Launch attendees")
        cls.audience.contacts.add(
            Contact.objects.create(workspace=cls.workspace, wa_id="27726124698",
                                   display_name="Thabo Mokoena")
        )

    def test_the_page_lists_templates_and_audiences(self):
        response = self.client.get(reverse("library:bulk_send"))
        self.assertContains(response, "logistics_update")

    def test_an_agent_may_not_open_the_bulk_send_page(self):
        self.client.force_login(self.agent)
        self.assertEqual(self.client.get(reverse("library:bulk_send")).status_code, 403)

    def test_posting_queues_the_background_send(self):
        with mock.patch(
            "apps.channels_wa.messaging.meta_cloud.MetaCloudChannel.send_template",
            return_value=SendResult(ok=True, wamid="wamid.UI"),
        ) as send:
            response = self.client.post(reverse("library:bulk_send"), {
                "template": self.template.pk,
                "audiences": [self.audience.pk],
            })
        from apps.library.models import BulkSend

        batch = BulkSend.objects.get()
        self.assertRedirects(response, reverse("reporting:bulk_send", args=[batch.pk]))
        send.assert_called_once()
        message = Message.objects.get(kind=Message.Kind.TEMPLATE)
        self.assertEqual(message.payload["values"], ["Thabo"])

    def test_choosing_no_audience_is_a_plain_error_and_sends_nothing(self):
        with mock.patch(
            "apps.channels_wa.messaging.meta_cloud.MetaCloudChannel.send_template"
        ) as send:
            response = self.client.post(reverse("library:bulk_send"), {
                "template": self.template.pk,
            }, follow=True)
        send.assert_not_called()
        self.assertContains(response, "at least one audience")

    def test_a_media_template_demands_a_file_or_link(self):
        self.template.components = [
            {"type": "HEADER", "format": "DOCUMENT"},
            {"type": "BODY", "text": "Hi {{1}}"},
        ]
        self.template.save(update_fields=["components"])
        with mock.patch(
            "apps.channels_wa.messaging.meta_cloud.MetaCloudChannel.send_template"
        ) as send:
            response = self.client.post(reverse("library:bulk_send"), {
                "template": self.template.pk,
                "audiences": [self.audience.pk],
            }, follow=True)
        send.assert_not_called()
        self.assertContains(response, "attach a file or paste an https link")


class AddNewPersonTests(AudienceUiBase):
    def test_a_person_who_never_messaged_can_be_created_straight_into_an_audience(self):
        from apps.contacts.models import Audience, Contact

        audience = Audience.objects.create(workspace=self.workspace, name="Launch attendees")
        response = self.client.post(
            reverse("contacts:audience_add_new", args=[audience.pk]),
            {"name": "Thandi Nkosi", "mobile": "082 123 4567"},
        )
        self.assertRedirects(response, reverse("contacts:audience_detail", args=[audience.pk]))
        contact = Contact.objects.get(wa_id="27821234567")
        self.assertEqual(contact.display_name, "Thandi Nkosi")
        self.assertIn(contact, audience.contacts.all())

        # Same number again: no twin, just already-there.
        self.client.post(reverse("contacts:audience_add_new", args=[audience.pk]), {"name": "T", "mobile": "+27 82 123 4567"})
        self.assertEqual(Contact.objects.filter(wa_id="27821234567").count(), 1)

        # A bad number is refused with a plain message.
        self.client.post(reverse("contacts:audience_add_new", args=[audience.pk]), {"name": "X", "mobile": "12"})
        self.assertEqual(Contact.objects.count(), 1)

    def test_agents_cannot_add_people(self):
        from apps.contacts.models import Audience

        audience = Audience.objects.create(workspace=self.workspace, name="Launch attendees")
        self.client.force_login(self.agent)
        response = self.client.post(reverse("contacts:audience_add_new", args=[audience.pk]), {"name": "A", "mobile": "0821234567"})
        self.assertEqual(response.status_code, 403)
